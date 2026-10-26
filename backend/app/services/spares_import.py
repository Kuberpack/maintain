"""Read the plant spares workbook and upsert rows.

The workbook is one area per sheet. Title rows sit above a header
`Iteam | Item Group | Qty`. Blank item groups are stored as "".
Two lines with the same name and group in one sheet are the same item;
their quantities are added so a re-import does not invent a second row
and does not drop stock.
"""

from __future__ import annotations

import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from app.models import Machine, SparePart

# Sheet area -> one machine name in the plant roster. Anything else stays
# unlinked: "Auto Die" and "Manul Die" each match several die cutters, and
# bearing descriptions often name more than one machine.
OBVIOUS_MACHINE_NAMES: dict[str, str] = {
    "flute laminetar": "Flute Laminator",
    "flute laminator": "Flute Laminator",
    "manual stitching": "Manual Stitching Machine (6 nos)",
}

_QTY_RE = re.compile(r"^([+-]?\d+(?:\.\d+)?)\s*(.*)$")
_HEADER_NAMES = {"iteam", "item"}


@dataclass(frozen=True)
class ParsedSpare:
    category: str
    name: str
    item_group: str
    qty_on_hand: float
    unit: str | None


@dataclass
class SheetStats:
    category: str
    sheet_rows: int = 0
    items: int = 0
    blank_rows_skipped: int = 0
    duplicate_rows_merged: int = 0
    bad_qty_skipped: int = 0


@dataclass
class ParseResult:
    items: list[ParsedSpare]
    sheets: list[SheetStats]

    @property
    def sheet_rows(self) -> int:
        return sum(sheet.sheet_rows for sheet in self.sheets)

    @property
    def duplicate_rows_merged(self) -> int:
        return sum(sheet.duplicate_rows_merged for sheet in self.sheets)

    @property
    def blank_rows_skipped(self) -> int:
        return sum(sheet.blank_rows_skipped for sheet in self.sheets)


@dataclass
class ImportStats:
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    linked: dict[str, str] = field(default_factory=dict)
    missing_machines: dict[str, str] = field(default_factory=dict)


def obvious_machine_name(category: str) -> str | None:
    return OBVIOUS_MACHINE_NAMES.get(category.strip().casefold())


def parse_qty(value: object) -> tuple[float, str | None] | None:
    """Return (number, unit). Unit is set when the cell was '12 Mtr' or '2 pkt'."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        qty = round(float(value), 3)
        if qty < 0:
            return None
        return (qty, None)
    if isinstance(value, str):
        match = _QTY_RE.match(value.strip())
        if match is None:
            return None
        qty = round(float(match.group(1)), 3)
        if qty < 0:
            return None
        unit = match.group(2).strip() or None
        if unit is not None and len(unit) > 32:
            return None
        return (qty, unit)
    return None


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_header(name: str, group: str) -> bool:
    return name.casefold() in _HEADER_NAMES and group.casefold() == "item group"


def parse_workbook(path: Path) -> ParseResult:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        items: list[ParsedSpare] = []
        sheets: list[SheetStats] = []
        for sheet_name in workbook.sheetnames:
            category = sheet_name.strip()
            stats = SheetStats(category=category)
            parsed = _parse_sheet(workbook[sheet_name].iter_rows(max_col=3, values_only=True), category, stats)
            items.extend(parsed)
            stats.items = len(parsed)
            sheets.append(stats)
        return ParseResult(items=items, sheets=sheets)
    finally:
        workbook.close()


def _parse_sheet(
    rows: Iterable[tuple[object, ...] | None],
    category: str,
    stats: SheetStats,
) -> list[ParsedSpare]:
    header_seen = False
    merged: dict[tuple[str, str], ParsedSpare] = {}
    order: list[tuple[str, str]] = []
    for row in rows:
        values = row or ()
        name = _cell_text(values[0] if len(values) > 0 else None)
        group = _cell_text(values[1] if len(values) > 1 else None)
        raw_qty = values[2] if len(values) > 2 else None
        if not header_seen:
            if _is_header(name, group):
                header_seen = True
            continue
        if _is_header(name, group):
            continue
        if not name and not group and (raw_qty is None or _cell_text(raw_qty) == ""):
            continue
        if not name:
            stats.blank_rows_skipped += 1
            continue
        parsed_qty = parse_qty(raw_qty)
        if parsed_qty is None:
            stats.bad_qty_skipped += 1
            continue
        qty, unit = parsed_qty
        stats.sheet_rows += 1
        key = (name, group)
        existing = merged.get(key)
        if existing is None:
            merged[key] = ParsedSpare(
                category=category,
                name=name,
                item_group=group,
                qty_on_hand=qty,
                unit=unit,
            )
            order.append(key)
            continue
        stats.duplicate_rows_merged += 1
        merged[key] = ParsedSpare(
            category=category,
            name=name,
            item_group=group,
            qty_on_hand=round(existing.qty_on_hand + qty, 3),
            unit=existing.unit or unit,
        )
    return [merged[key] for key in order]


def _same_qty(left: float, right: float) -> bool:
    return round(float(left), 3) == round(float(right), 3)


def apply_import(db: Session, items: list[ParsedSpare]) -> ImportStats:
    """Upsert by category + name + item group. Does not clear min_qty.

    A machine link is written only when the sheet area maps to exactly one
    machine. A link set by hand on any other area is left as it is. Re-running
    sets on-hand qty back to the sheet.
    """
    by_name: dict[str, list[Machine]] = {}
    for machine in db.query(Machine).all():
        by_name.setdefault(machine.name.casefold(), []).append(machine)

    stats = ImportStats()
    resolved_machine: dict[str, object] = {}

    def machine_id_for(category: str):
        if category in resolved_machine:
            return resolved_machine[category]
        target = obvious_machine_name(category)
        if target is None:
            resolved_machine[category] = None
            return None
        matches = by_name.get(target.casefold(), [])
        if len(matches) == 1:
            resolved_machine[category] = matches[0].id
            stats.linked[category] = matches[0].name
            return matches[0].id
        resolved_machine[category] = None
        stats.missing_machines[category] = target
        return None

    for item in items:
        machine_id = machine_id_for(item.category)
        existing = (
            db.query(SparePart)
            .filter(
                SparePart.category == item.category,
                SparePart.name == item.name,
                SparePart.item_group == item.item_group,
            )
            .one_or_none()
        )
        if existing is None:
            db.add(
                SparePart(
                    name=item.name,
                    item_group=item.item_group,
                    category=item.category,
                    qty_on_hand=item.qty_on_hand,
                    unit=item.unit,
                    machine_id=machine_id,
                )
            )
            stats.created += 1
            continue
        changed = False
        if not _same_qty(existing.qty_on_hand, item.qty_on_hand):
            existing.qty_on_hand = item.qty_on_hand
            changed = True
        if (existing.unit or None) != (item.unit or None):
            existing.unit = item.unit
            changed = True
        if machine_id is not None and existing.machine_id != machine_id:
            existing.machine_id = machine_id
            changed = True
        if changed:
            stats.updated += 1
        else:
            stats.unchanged += 1
    db.commit()
    return stats
